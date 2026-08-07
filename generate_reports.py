import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def build_weekly_report(doc, data):
    # Header Title
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("CHAROTAR UNIVERSITY OF SCIENCE & TECHNOLOGY\nDEVANG PATEL INSTITUTE OF ADVANCE TECHNOLOGY AND RESEARCH\nDEPSTAR")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(12)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0, 51, 102)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("WEEKLY REPORT")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.underline = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("In this report students require to mention the work has been done during the week and planning for the next week. (Please fill individual report for each week)")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(60, 60, 60)

    # Metadata Table
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="000000", sz="6")

    meta_items = [
        [("Project ID:", data["project_id"]), ("Student ID:", data["student_id"])],
        [("From Date:", data["from_date"]), ("To Date:", data["to_date"])],
        [("Semester:", data["semester"]), ("Internship ID:", data["internship_id"])]
    ]

    for row_idx, row_data in enumerate(meta_items):
        for col_idx, (label, val) in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            cell.width = Inches(3.2)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            r_lbl = p.add_run(f"{label} ")
            r_lbl.font.name = "Arial"
            r_lbl.font.size = Pt(10)
            r_lbl.font.bold = True
            
            r_val = p.add_run(val)
            r_val.font.name = "Arial"
            r_val.font.size = Pt(10)
            r_val.font.bold = False

    # Section 1: Work done
    p_work_head = doc.add_paragraph()
    p_work_head.paragraph_format.space_before = Pt(16)
    p_work_head.paragraph_format.space_after = Pt(4)
    r_wh1 = p_work_head.add_run(f"Work done from Date: {data['from_date']} to {data['to_date']}\n")
    r_wh1.font.name = "Arial"
    r_wh1.font.size = Pt(11)
    r_wh1.font.bold = True

    r_wh2 = p_work_head.add_run("(Attach supporting Documents):")
    r_wh2.font.name = "Arial"
    r_wh2.font.size = Pt(10.5)
    r_wh2.font.bold = True

    # Work done content box
    tbl_work = doc.add_table(rows=1, cols=1)
    tbl_work.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_work, color="000000", sz="6")
    cell_work = tbl_work.cell(0, 0)
    cell_work.width = Inches(6.5)
    set_cell_margins(cell_work, top=120, bottom=120, left=150, right=150)
    set_cell_background(cell_work, "FAFAFA")

    # Add work items
    p_w_init = cell_work.paragraphs[0]
    p_w_init.paragraph_format.space_after = Pt(4)
    r_summary = p_w_init.add_run(f"Module Focus: {data['work_summary']}\n")
    r_summary.font.name = "Arial"
    r_summary.font.size = Pt(10)
    r_summary.font.bold = True

    for bullet in data["work_bullets"]:
        p_b = cell_work.add_paragraph()
        p_b.paragraph_format.space_before = Pt(2)
        p_b.paragraph_format.space_after = Pt(3)
        p_b.paragraph_format.left_indent = Inches(0.2)
        r_bullet = p_b.add_run(f"• {bullet}")
        r_bullet.font.name = "Arial"
        r_bullet.font.size = Pt(9.5)

    # Section 2: Plans for next week
    p_plan_head = doc.add_paragraph()
    p_plan_head.paragraph_format.space_before = Pt(14)
    p_plan_head.paragraph_format.space_after = Pt(4)
    r_ph = p_plan_head.add_run(f"Plans for next week: Date: {data['next_from_date']} to {data['next_to_date']}")
    r_ph.font.name = "Arial"
    r_ph.font.size = Pt(11)
    r_ph.font.bold = True

    # Plans content box
    tbl_plan = doc.add_table(rows=1, cols=1)
    tbl_plan.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_plan, color="000000", sz="6")
    cell_plan = tbl_plan.cell(0, 0)
    cell_plan.width = Inches(6.5)
    set_cell_margins(cell_plan, top=120, bottom=120, left=150, right=150)
    set_cell_background(cell_plan, "FAFAFA")

    p_p_init = cell_plan.paragraphs[0]
    p_p_init.paragraph_format.space_after = Pt(4)
    r_plan_sum = p_p_init.add_run(f"Target Objectives:\n")
    r_plan_sum.font.name = "Arial"
    r_plan_sum.font.size = Pt(10)
    r_plan_sum.font.bold = True

    for p_bullet in data["plan_bullets"]:
        p_pb = cell_plan.add_paragraph()
        p_pb.paragraph_format.space_before = Pt(2)
        p_pb.paragraph_format.space_after = Pt(3)
        p_pb.paragraph_format.left_indent = Inches(0.2)
        r_pbullet = p_pb.add_run(f"• {p_bullet}")
        r_pbullet.font.name = "Arial"
        r_pbullet.font.size = Pt(9.5)

    # Section 3: References
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_before = Pt(14)
    p_ref.paragraph_format.space_after = Pt(4)
    r_ref_h = p_ref.add_run("References:\n")
    r_ref_h.font.name = "Arial"
    r_ref_h.font.size = Pt(11)
    r_ref_h.font.bold = True
    r_ref_h.font.underline = True

    for ref in data["references"]:
        p_rf = doc.add_paragraph()
        p_rf.paragraph_format.space_before = Pt(1)
        p_rf.paragraph_format.space_after = Pt(2)
        p_rf.paragraph_format.left_indent = Inches(0.15)
        r_rf = p_rf.add_run(f"1. {ref}" if data["references"].index(ref)==0 else f"2. {ref}")
        r_rf.font.name = "Arial"
        r_rf.font.size = Pt(9.5)

    # Signatures Footer
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(30)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_sig1 = p_sig.add_run("Signature of Student                                                          Signature of Mentor")
    r_sig1.font.name = "Arial"
    r_sig1.font.size = Pt(10.5)
    r_sig1.font.bold = True

# Data for 3 weeks
week1_data = {
    "project_id": "PRJ_IT_5_2026_10",
    "student_id": "D25DIT083",
    "from_date": "18/07/2026",
    "to_date": "24/07/2026",
    "next_from_date": "25/07/2026",
    "next_to_date": "31/07/2026",
    "semester": "5th",
    "internship_id": "INT_2026_D25DIT083",
    "work_summary": "System Architecture, Authentication Module & Database Schema Definition",
    "work_bullets": [
        "Project Architecture & Environment Setup: Configured React 19 + Vite frontend SPA and Node.js + Express backend server with CORS & cookie-parser.",
        "Database Schemas Definition: Created Mongoose ODM models for User (auth & profile), Expense (transactions), Budget (category limits), Friend (social), Message (P2P chat), and Group (bill splits).",
        "User Registration & Login API: Implemented POST /api/auth/register and POST /api/auth/login with password hashing via bcryptjs and 7-day session JWT tokens.",
        "Security PIN Lock & Recovery Flow: Implemented SetPin.jsx, PinLock.jsx views, password constraints validation, and reset password token generation.",
        "Google OAuth Integration: Integrated GoogleOAuthProvider and GoogleLogin component setup for smooth authentication."
    ],
    "plan_bullets": [
        "Build core personal financial ledger endpoints (Income tracking, Expense management with category filtering).",
        "Implement Category Budget limit configuration and visual progress alert indicators.",
        "Develop Summary Analytics API (GET /summary, GET /monthly, GET /by-category) and frontend chart widgets."
    ],
    "references": [
        "GitHub Repository: https://github.com/harshil5506/Trackify.git (Branch: main, auth-integration)",
        "React.js & Node.js Express Documentation | Mongoose ODM & JWT Authentication Standards"
    ]
}

week2_data = {
    "project_id": "PRJ_IT_5_2026_10",
    "student_id": "D25DIT083",
    "from_date": "25/07/2026",
    "to_date": "31/07/2026",
    "next_from_date": "01/08/2026",
    "next_to_date": "07/08/2026",
    "semester": "5th",
    "internship_id": "INT_2026_D25DIT083",
    "work_summary": "Personal Financial Tracking, Category Budgets & Multi-Format Data Export Engine",
    "work_bullets": [
        "Income & Expense Ledger: Developed full CRUD backend API (expenses.js) supporting transaction title, category, payment method, date range, and pagination limits.",
        "Category Budgeting System: Built budget.js CRUD routes and Budget.jsx view to configure monthly category limits and render active overrun warnings.",
        "Summary Analytics Engine: Developed analytics.js endpoints (GET /summary, GET /monthly comparison, GET /by-category aggregation) and integrated Recharts widgets.",
        "Multi-Format Data Export Engine: Created csvExport.js utility and enhanced Reports.jsx & Transactions.jsx to support PDF (jsPDF + autotable), CSV, and JSON report downloads.",
        "Data Seeding & Verification: Built seed_trackify.js script populating initial user financial ledgers and verified client production build."
    ],
    "plan_bullets": [
        "Develop Social & Group Expense Sharing module (Friends list, Direct peer messaging, Group bill creation).",
        "Build flexible Bill Splitting engine supporting Equal, Custom amount, and Percentage-based allocations.",
        "Develop Activity Heatmap grid and Spending Personality Quiz module to reach Checkpoint 1 milestone."
    ],
    "references": [
        "GitHub Repository: https://github.com/harshil5506/Trackify.git",
        "jsPDF & jspdf-autotable Documentation | Recharts Visualization API Guide"
    ]
}

week3_data = {
    "project_id": "PRJ_IT_5_2026_10",
    "student_id": "D25DIT083",
    "from_date": "01/08/2026",
    "to_date": "07/08/2026",
    "next_from_date": "08/08/2026",
    "next_to_date": "14/08/2026",
    "semester": "5th",
    "internship_id": "INT_2026_D25DIT083",
    "work_summary": "Group Bill Splitting, Peer Chat, Activity Heatmap, Spending Quiz & Checkpoint 1 Milestone",
    "work_bullets": [
        "Group Expense & Bill Splitting Engine: Built groups.js backend routes and GroupDetail.jsx frontend supporting Equal, Custom amount, and Percentage allocation splits with full/partial debt settlement.",
        "Peer-to-Peer Social Messaging: Implemented Friend.js & Message.js models, friends.js & messages.js API routes, and real-time chat interface in Friends.jsx.",
        "Activity Heatmap Component: Built ActivityHeatmap.jsx and GET /api/analytics/heatmap endpoint rendering 365-day contribution grid with intensity levels, tooltips, and 90/180/365-day range controls.",
        "Spending Personality Quiz: Created Quiz.jsx assessment with 5 behavioral questions, archetype score calculation (Master Saver 🛡️, Balanced Budgeter 📈, Experience Enthusiast 🛍️, Impulse Adventurer ⚡), and MongoDB profile persistence.",
        "Profile Page Integration & Checkpoint 1: Transferred Activity Heatmap and Spending Quiz badge to Profile.jsx, executed client build, and tagged release checkpoint-1 on GitHub."
    ],
    "plan_bullets": [
        "Perform end-to-end user workflow testing, accessibility audits, and response optimization across devices.",
        "Prepare final project documentation, demonstration recordings, and staging deployment."
    ],
    "references": [
        "GitHub Repository: https://github.com/harshil5506/Trackify.git (Tag: checkpoint-1, Commit: be0f2af)",
        "Trackify Project Walkthrough & Implementation Specification Documents"
    ]
}

# 1. Create individual week docx files
for idx, wdata in enumerate([week1_data, week2_data, week3_data], 1):
    doc = docx.Document()
    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
    
    build_weekly_report(doc, wdata)
    filename = f"Trackify_Weekly_Report_Week_{idx}.docx"
    doc.save(filename)
    print(f"Created {filename}")

# 2. Create combined docx file containing all 3 weeks
doc_combined = docx.Document()
for section in doc_combined.sections:
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

for idx, wdata in enumerate([week1_data, week2_data, week3_data], 1):
    if idx > 1:
        doc_combined.add_page_break()
    build_weekly_report(doc_combined, wdata)

combined_filename = "Trackify_Weekly_Reports_Week_1_to_3.docx"
doc_combined.save(combined_filename)
print(f"Created {combined_filename}")

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation_pptx():
    prs = Presentation()
    
    # 16:9 Widescreen aspect ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Tokens
    COLOR_BG_DARK = RGBColor(11, 15, 25)      # #0b0f19
    COLOR_CARD_BG = RGBColor(22, 30, 49)      # #161e31
    COLOR_PRIMARY = RGBColor(26, 46, 168)     # #1a2ea8
    COLOR_SECONDARY = RGBColor(45, 71, 201)   # #2d47c9
    COLOR_CYAN = RGBColor(6, 182, 212)        # #06b6d4
    COLOR_PURPLE = RGBColor(139, 92, 246)     # #8b5cf6
    COLOR_GREEN = RGBColor(16, 185, 129)      # #10b981
    COLOR_TEXT_WHITE = RGBColor(255, 255, 255)
    COLOR_TEXT_MUTED = RGBColor(156, 163, 175) # #9ca3af

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, slide_num, title_text, category_text="TRACKIFY"):
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf = badge_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{category_text}  |  SLIDE {slide_num}"
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_CYAN

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title_text
        r_t.font.size = Pt(24)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = RGBColor(50, 60, 85)
        shape.line.width = Pt(1)
        return shape

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, COLOR_BG_DARK)

    add_card(s1, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1), RGBColor(18, 24, 40))

    tb1 = s1.shapes.add_textbox(Inches(1.6), Inches(1.6), Inches(10.1), Inches(4.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    r = p.add_run()
    r.text = "PROJECT PROGRESS PRESENTATION (WEEKS 1 TO 3)\n"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = COLOR_CYAN

    p2 = tf1.add_paragraph()
    r2 = p2.add_run()
    r2.text = "TRACKIFY"
    r2.font.size = Pt(40)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_TEXT_WHITE

    p3 = tf1.add_paragraph()
    r3 = p3.add_run()
    r3.text = "Full-Stack Personal Finance & Group Expense Sharing Platform\n"
    r3.font.size = Pt(18)
    r3.font.color.rgb = COLOR_TEXT_MUTED

    p4 = tf1.add_paragraph()
    r4 = p4.add_run()
    r4.text = "----------------------------------------------------------------------------------------------------\n"
    r4.font.size = Pt(10)
    r4.font.color.rgb = RGBColor(60, 75, 105)

    p5 = tf1.add_paragraph()
    r5 = p5.add_run()
    r5.text = "• Institution: CHARUSAT - Devang Patel Institute of Advance Technology & Research (DEPSTAR)\n"
    r5.text += "• Semester: 5th  |  Project ID: PRJ_IT_5_2026_10  |  Domain: Web Development / FinTech\n"
    r5.text += "• Student ID: D25DIT083  |  Internship ID: INT_2026_D25DIT083\n"
    r5.text += "• Milestone: Checkpoint 1 Released (checkpoint-1)  |  Repository: github.com/harshil5506/Trackify"
    r5.font.size = Pt(13)
    r5.font.color.rgb = COLOR_TEXT_WHITE

    # ==========================================
    # SLIDE 2: 3-WEEK EXECUTIVE OVERVIEW
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, COLOR_BG_DARK)
    add_header(s2, 2, "3-Week Milestone Overview (Weeks 1 to 3)")

    col_width = Inches(3.64)
    card_h = Inches(5.4)

    # Week 1 Card
    add_card(s2, Inches(0.8), Inches(1.5), col_width, card_h)
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(3.24), card_h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "WEEK 1 (18/07 - 24/07)\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_CYAN

    p = tf.add_paragraph()
    r1 = p.add_run()
    r1.text = "System Architecture & Security\n\n"
    r1.font.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = COLOR_TEXT_WHITE

    r2 = p.add_run()
    r2.text = "• React 19 + Vite SPA frontend & Node.js Express backend setup.\n• Defined 6 Mongoose schemas (User, Expense, Budget, Friend, Message, Group).\n• JWT session auth & bcryptjs password hashing.\n• Security PIN lock views (SetPin & PinLock).\n• Google OAuth login integration."
    r2.font.size = Pt(11)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # Week 2 Card
    add_card(s2, Inches(4.84), Inches(1.5), col_width, card_h)
    tb = s2.shapes.add_textbox(Inches(5.04), Inches(1.7), Inches(3.24), card_h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "WEEK 2 (25/07 - 31/07)\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_PURPLE

    p = tf.add_paragraph()
    r1 = p.add_run()
    r1.text = "Ledger, Budgets & Reports\n\n"
    r1.font.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = COLOR_TEXT_WHITE

    r2 = p.add_run()
    r2.text = "• Income & Expense CRUD API endpoints.\n• Monthly Category Budget limits & overrun warning alerts.\n• Aggregation pipelines & Recharts visualization widgets.\n• Multi-format export engine (PDF, CSV, JSON).\n• Seeded 75+ test transactions."
    r2.font.size = Pt(11)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # Week 3 Card
    add_card(s2, Inches(8.88), Inches(1.5), col_width, card_h)
    tb = s2.shapes.add_textbox(Inches(9.08), Inches(1.7), Inches(3.24), card_h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "WEEK 3 (01/08 - 07/08)\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_GREEN

    p = tf.add_paragraph()
    r1 = p.add_run()
    r1.text = "Multi-Currency & Email Automation\n\n"
    r1.font.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = COLOR_TEXT_WHITE

    r2 = p.add_run()
    r2.text = "• Multi-currency preference selection (INR, USD, EUR, GBP).\n• Dynamic currency formatting across ledgers & exports.\n• Automated monthly report emails via Nodemailer.\n• Verified production build (npm run build).\n• Checkpoint 1 milestone tag."
    r2.font.size = Pt(11)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # ==========================================
    # SLIDE 3: WEEK 1 DETAILS
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, COLOR_BG_DARK)
    add_header(s3, 3, "Week 1: System Architecture & Authentication Module")

    add_card(s3, Inches(0.8), Inches(1.5), Inches(7.5), Inches(5.4))
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(7.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Key Accomplishments & Deliverables:\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_CYAN

    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "1. MERN + Vite Architecture Setup:\n   Configured React 19 SPA frontend with Vite HMR and Express backend server.\n\n2. Database Schema Definition (Mongoose ODM):\n   Designed User, Expense, Budget, Friend, Message, and Group schemas.\n\n3. Session Security & Hashing:\n   Implemented bcryptjs password hashing and 7-day JWT session tokens.\n\n4. Security PIN Lock Flow:\n   Built SetPin.jsx & PinLock.jsx to guard sensitive financial ledgers."
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # Right Card Stats
    add_card(s3, Inches(8.6), Inches(1.5), Inches(3.933), Inches(5.4))
    tb_r = s3.shapes.add_textbox(Inches(8.8), Inches(1.8), Inches(3.533), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    
    pr = tf_r.paragraphs[0]
    r = pr.add_run()
    r.text = "MODELS CREATED\n"
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_TEXT_MUTED
    
    pr2 = tf_r.add_paragraph()
    r = pr2.add_run()
    r.text = "6 Data Schemas\n\n"
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = COLOR_CYAN

    pr3 = tf_r.add_paragraph()
    r = pr3.add_run()
    r.text = "AUTH ENDPOINTS:\n• POST /api/auth/register\n• POST /api/auth/login\n• POST /api/auth/google\n• POST /api/auth/set-pin\n• POST /api/auth/verify-pin"
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_TEXT_WHITE

    # ==========================================
    # SLIDE 4: WEEK 2 DETAILS
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, COLOR_BG_DARK)
    add_header(s4, 4, "Week 2: Financial Ledger, Budgets & Multi-Format Exports")

    add_card(s4, Inches(0.8), Inches(1.5), Inches(7.5), Inches(5.4))
    tb = s4.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(7.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Key Accomplishments & Deliverables:\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_PURPLE

    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "1. Personal Financial Ledger CRUD:\n   Developed expenses.js endpoints for transactions, category tags, payment methods, date ranges, and pagination.\n\n2. Category Budgeting System:\n   Created budget.js routes and Budget.jsx view with visual overrun progress bars.\n\n3. Analytics & Recharts Widgets:\n   Built analytics.js endpoints (GET /summary, GET /monthly, GET /by-category).\n\n4. Multi-Format Data Export Engine:\n   Built csvExport.js supporting PDF (jsPDF + autotable), CSV, and JSON downloads."
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # Right Card Stats
    add_card(s4, Inches(8.6), Inches(1.5), Inches(3.933), Inches(5.4))
    tb_r = s4.shapes.add_textbox(Inches(8.8), Inches(1.8), Inches(3.533), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    
    pr = tf_r.paragraphs[0]
    r = pr.add_run()
    r.text = "EXPORT FORMATS\n"
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_TEXT_MUTED
    
    pr2 = tf_r.add_paragraph()
    r = pr2.add_run()
    r.text = "PDF, CSV, JSON\n\n"
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = COLOR_PURPLE

    pr3 = tf_r.add_paragraph()
    r = pr3.add_run()
    r.text = "DATABASE SEEDING:\n• Generated seed_trackify.js script\n• Populated 75+ sample expenses\n• Seeded roommate group splits & direct messaging data"
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_TEXT_WHITE

    # ==========================================
    # SLIDE 5: WEEK 3 DETAILS
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, COLOR_BG_DARK)
    add_header(s5, 5, "Week 3: Multi-Currency Support & Automated Email Reports")

    add_card(s5, Inches(0.8), Inches(1.5), Inches(7.5), Inches(5.4))
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(7.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Key Accomplishments & Deliverables:\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_GREEN

    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "1. Multi-Currency Preference & Formatting:\n   Added currency selection (INR ₹, USD $, EUR €, GBP £) in Profile.jsx and applied dynamic currency formatting across transaction tables, summaries, and exports.\n\n2. Automated Email Reports (Nodemailer):\n   Integrated Nodemailer email transport to automatically deliver formatted monthly summaries directly to user email accounts.\n\n3. Checkpoint 1 Milestone Release:\n   Executed production build (npm run build) and tagged release 'checkpoint-1' on GitHub main branch."
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # Right Card Stats
    add_card(s5, Inches(8.6), Inches(1.5), Inches(3.933), Inches(5.4))
    tb_r = s5.shapes.add_textbox(Inches(8.8), Inches(1.8), Inches(3.533), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    
    pr = tf_r.paragraphs[0]
    r = pr.add_run()
    r.text = "MILESTONE STATUS\n"
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_TEXT_MUTED
    
    pr2 = tf_r.add_paragraph()
    r = pr2.add_run()
    r.text = "v1.0-cp1\n\n"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = COLOR_GREEN

    pr3 = tf_r.add_paragraph()
    r = pr3.add_run()
    r.text = "CHECKPOINT 1 TAGGED:\n• Tag: checkpoint-1\n• Commit: be0f2af\n• Verified production build\n• Automated Nodemailer dispatch"
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_TEXT_WHITE

    # ==========================================
    # SLIDE 6: TECH STACK SPECIFICATIONS
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, COLOR_BG_DARK)
    add_header(s6, 6, "Full-Stack Technology Architecture")

    c_w = Inches(3.64)
    c_h = Inches(5.4)

    # Frontend Card
    add_card(s6, Inches(0.8), Inches(1.5), c_w, c_h)
    tb = s6.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(3.24), c_h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "FRONTEND LAYER\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_CYAN
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "• React 19 Single Page App\n• Vite Build System & HMR\n• Recharts Interactive Charts\n• jsPDF + AutoTable Export\n• Custom CSS Design System"
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # Backend Card
    add_card(s6, Inches(4.84), Inches(1.5), c_w, c_h)
    tb = s6.shapes.add_textbox(Inches(5.04), Inches(1.7), Inches(3.24), c_h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "BACKEND LAYER\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_PURPLE
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "• Node.js & Express REST API\n• JWT Token Authentication\n• Bcryptjs Password Hashing\n• Nodemailer Mail Dispatch\n• CORS & Cookie Parser"
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # Database & DevOps Card
    add_card(s6, Inches(8.88), Inches(1.5), c_w, c_h)
    tb = s6.shapes.add_textbox(Inches(9.08), Inches(1.7), Inches(3.24), c_h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "DATABASE & DEVOPS\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_GREEN
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "• MongoDB NoSQL Database\n• Mongoose ODM Schema Modeling\n• Git & GitHub Versioning\n• Seed Data Generators\n• Checkpoint Release Tagging"
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # ==========================================
    # SLIDE 7: FUTURE ROADMAP
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, COLOR_BG_DARK)
    add_header(s7, 7, "Future Roadmap & Upcoming Objectives (Weeks 4+)")

    # Left Card
    add_card(s7, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.4))
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.2), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "GROUP EXPENSE DEBT SETTLEMENT\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_PURPLE
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "• Extend multi-currency support to shared group ledgers.\n• Implement debt minimization algorithms to settle roommate expenses with minimal transactions.\n• Add real-time expense notifications & group P2P messaging."
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # Right Card
    add_card(s7, Inches(6.8), Inches(1.5), Inches(5.733), Inches(5.4))
    tb = s7.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.333), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "ADVANCED FINANCIAL INSIGHTS\n"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_CYAN
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "• Automated recurring subscription detection & bill due reminders.\n• Smart spending habit predictions and monthly budget suggestions.\n• Mobile responsive touch gesture controls for quick logging."
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_TEXT_WHITE

    # ==========================================
    # SLIDE 8: CONCLUSION SLIDE
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, COLOR_BG_DARK)

    add_card(s8, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5), RGBColor(18, 24, 40))
    tb = s8.shapes.add_textbox(Inches(1.8), Inches(2.0), Inches(9.733), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "THANK YOU!\n"
    r.font.size = Pt(36)
    r.font.bold = True
    r.font.color.rgb = COLOR_GREEN

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Trackify Weeks 1–3 Development & Checkpoint 1 Milestone Reached 🎉\n\n"
    r2.font.size = Pt(16)
    r2.font.color.rgb = COLOR_TEXT_MUTED

    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = "GitHub Repository: github.com/harshil5506/Trackify (Tag: checkpoint-1)"
    r3.font.size = Pt(14)
    r3.font.bold = True
    r3.font.color.rgb = COLOR_TEXT_WHITE

    # Save PPTX Presentation
    prs.save("Trackify_Weekly_Reports_Presentation.pptx")
    print("SUCCESS: Trackify_Weekly_Reports_Presentation.pptx created successfully!")

if __name__ == "__main__":
    create_presentation_pptx()
